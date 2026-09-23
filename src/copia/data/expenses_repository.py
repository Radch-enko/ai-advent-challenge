from __future__ import annotations

import json
import os
import tempfile
from datetime import UTC, datetime
from pathlib import Path
from uuid import uuid4
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.styles import Font
from openpyxl.utils.exceptions import InvalidFileException
from pydantic import ValidationError

from ..domain.models.expense import Expense, ExpenseCreate

SHEET_NAME = "Expenses"
HEADERS = (
    "ID",
    "Occurred At",
    "Name",
    "Category",
    "Amount (RUB)",
    "Merchant",
    "Payment Method",
    "Note",
    "Tags",
    "Created At",
)


class ExpensesFileNotFoundError(FileNotFoundError):
    pass


class ExpensesStorageError(RuntimeError):
    pass


class ExpensesRepository:
    """Reads and appends expenses in a user-managed Excel workbook."""

    def __init__(self, path: Path) -> None:
        self._path = path.expanduser()

    def list(self) -> list[Expense]:
        workbook = self._load_workbook()
        try:
            worksheet = self._validated_worksheet(workbook)
            expenses = [
                self._parse_row(row_number, values)
                for row_number, values in enumerate(
                    worksheet.iter_rows(min_row=2, max_col=len(HEADERS), values_only=True),
                    start=2,
                )
                if any(value is not None for value in values)
            ]
            return sorted(expenses, key=lambda expense: expense.occurred_at, reverse=True)
        finally:
            workbook.close()

    def add(self, expense_create: ExpenseCreate) -> Expense:
        workbook = self._load_workbook()
        temporary_path: Path | None = None
        try:
            worksheet = self._validated_worksheet(workbook)
            for row_number, values in enumerate(
                worksheet.iter_rows(min_row=2, max_col=len(HEADERS), values_only=True),
                start=2,
            ):
                if any(value is not None for value in values):
                    self._parse_row(row_number, values)

            expense = Expense(
                **expense_create.model_dump(),
                id=uuid4(),
                created_at=datetime.now(UTC),
            )
            worksheet.append(self._serialize(expense))
            self._format_worksheet(worksheet)

            with tempfile.NamedTemporaryFile(
                dir=self._path.parent,
                prefix=".finances-",
                suffix=".xlsx",
                delete=False,
            ) as temporary:
                temporary_path = Path(temporary.name)
            workbook.save(temporary_path)
            os.replace(temporary_path, self._path)
            temporary_path = None
            return expense
        except ExpensesStorageError:
            raise
        except (OSError, ValueError) as error:
            raise ExpensesStorageError(f"Unable to save workbook: {error}") from error
        finally:
            workbook.close()
            if temporary_path is not None:
                temporary_path.unlink(missing_ok=True)

    def _load_workbook(self):
        if not self._path.is_file():
            raise ExpensesFileNotFoundError(f"Finances workbook not found: {self._path}")
        try:
            return load_workbook(self._path)
        except (BadZipFile, InvalidFileException, OSError, ValueError) as error:
            raise ExpensesStorageError(f"Unable to open finances workbook: {error}") from error

    @staticmethod
    def _validated_worksheet(workbook):
        if SHEET_NAME not in workbook.sheetnames:
            raise ExpensesStorageError(f"Workbook must contain the '{SHEET_NAME}' sheet")
        worksheet = workbook[SHEET_NAME]
        actual_headers = tuple(
            worksheet.cell(row=1, column=column).value for column in range(1, len(HEADERS) + 1)
        )
        if actual_headers != HEADERS:
            expected = ", ".join(HEADERS)
            raise ExpensesStorageError(f"Invalid workbook headers; expected: {expected}")
        return worksheet

    @staticmethod
    def _parse_row(row_number: int, values: tuple[object, ...]) -> Expense:
        (
            expense_id,
            occurred_at,
            name,
            category,
            amount_rub,
            merchant,
            payment_method,
            note,
            tags_json,
            created_at,
        ) = values
        if not isinstance(tags_json, str):
            raise ExpensesStorageError(
                f"Invalid Tags value in row {row_number}: expected JSON array"
            )
        try:
            tags = json.loads(tags_json)
        except json.JSONDecodeError as error:
            raise ExpensesStorageError(
                f"Invalid Tags JSON in row {row_number}: {error.msg}"
            ) from error
        if not isinstance(tags, list) or not all(isinstance(tag, str) for tag in tags):
            raise ExpensesStorageError(
                f"Invalid Tags value in row {row_number}: expected JSON array"
            )
        try:
            return Expense.model_validate(
                {
                    "id": expense_id,
                    "occurred_at": occurred_at,
                    "name": name,
                    "category": category,
                    "amount_rub": amount_rub,
                    "merchant": merchant,
                    "payment_method": payment_method,
                    "note": note,
                    "tags": tags,
                    "created_at": created_at,
                }
            )
        except ValidationError as error:
            reason = error.errors(include_url=False)[0]["msg"]
            raise ExpensesStorageError(
                f"Invalid expense data in row {row_number}: {reason}"
            ) from error

    @staticmethod
    def _serialize(expense: Expense) -> list[object]:
        return [
            str(expense.id),
            expense.occurred_at.isoformat(),
            expense.name,
            expense.category,
            float(expense.amount_rub),
            expense.merchant,
            expense.payment_method,
            expense.note,
            json.dumps(expense.tags, ensure_ascii=False),
            expense.created_at.isoformat(),
        ]

    @staticmethod
    def _format_worksheet(worksheet) -> None:
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = f"A1:J{worksheet.max_row}"
        widths = (38, 27, 28, 22, 16, 24, 20, 36, 28, 27)
        for column, width in enumerate(widths, start=1):
            worksheet.column_dimensions[worksheet.cell(1, column).column_letter].width = width
            worksheet.cell(1, column).font = Font(bold=True)
        row_number = worksheet.max_row
        worksheet.cell(row_number, 5).number_format = '#,##0.00 "RUB"'
        for column in (1, 2, 3, 4, 6, 7, 8, 9, 10):
            cell = worksheet.cell(row_number, column)
            if cell.value is not None:
                cell.data_type = "s"
