from datetime import UTC, datetime
from decimal import Decimal
from pathlib import Path
from uuid import UUID

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from copia.api import service
from copia.api.service import app
from copia.data import expenses_repository
from copia.data.expenses_repository import HEADERS, ExpensesRepository, ExpensesStorageError
from copia.domain.models.expense import ExpenseCreate


def _create_workbook(path: Path, rows: list[list[object]] | None = None) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Expenses"
    worksheet.append(list(HEADERS))
    for row in rows or []:
        worksheet.append(row)
    workbook.save(path)
    workbook.close()


def _row(
    expense_id: str,
    occurred_at: str,
    *,
    name: str = "Coffee",
    tags: str = '["food"]',
) -> list[object]:
    return [
        expense_id,
        occurred_at,
        name,
        "Food",
        250.5,
        "Cafe",
        "Card",
        "Morning",
        tags,
        "2026-09-23T08:00:00+00:00",
    ]


def test_repository_lists_newest_first_and_preserves_row_order_for_ties(tmp_path) -> None:
    path = tmp_path / "finances.xlsx"
    first_id = "00000000-0000-0000-0000-000000000001"
    second_id = "00000000-0000-0000-0000-000000000002"
    newest_id = "00000000-0000-0000-0000-000000000003"
    _create_workbook(
        path,
        [
            _row(first_id, "2026-09-22T10:00:00+06:00"),
            _row(second_id, "2026-09-22T10:00:00+06:00"),
            _row(newest_id, "2026-09-23T10:00:00+06:00"),
        ],
    )

    result = ExpensesRepository(path).list()

    assert [str(expense.id) for expense in result] == [newest_id, first_id, second_id]


def test_repository_appends_expense_and_applies_human_readable_formatting(tmp_path) -> None:
    path = tmp_path / "finances.xlsx"
    existing_id = "00000000-0000-0000-0000-000000000001"
    _create_workbook(path, [_row(existing_id, "2026-09-22T10:00:00+06:00")])
    repository = ExpensesRepository(path)

    created = repository.add(
        ExpenseCreate(
            occurred_at=datetime.fromisoformat("2026-09-23T14:30:00+06:00"),
            name="=SUM(A1:A2)",
            category="Food",
            amount_rub=Decimal("199.90"),
            merchant="Market",
            payment_method="Card",
            note="Weekly groceries",
            tags=["food", "weekly"],
        )
    )

    assert isinstance(created.id, UUID)
    assert created.created_at.tzinfo is UTC
    stored = repository.list()
    assert [str(expense.id) for expense in stored] == [str(created.id), existing_id]
    assert stored[0].amount_rub == Decimal("199.9")
    assert stored[0].tags == ["food", "weekly"]

    workbook = load_workbook(path, data_only=False)
    worksheet = workbook["Expenses"]
    assert worksheet.max_row == 3
    assert worksheet.freeze_panes == "A2"
    assert worksheet.auto_filter.ref == "A1:J3"
    assert worksheet["C3"].value == "=SUM(A1:A2)"
    assert worksheet["C3"].data_type == "s"
    assert worksheet["E3"].number_format == '#,##0.00 "RUB"'
    workbook.close()


def test_repository_rejects_invalid_headers_without_modifying_file(tmp_path) -> None:
    path = tmp_path / "finances.xlsx"
    _create_workbook(path)
    workbook = load_workbook(path)
    workbook["Expenses"]["A1"] = "Wrong"
    workbook.save(path)
    workbook.close()
    original = path.read_bytes()

    try:
        ExpensesRepository(path).add(
            ExpenseCreate(
                occurred_at=datetime.now(UTC),
                name="Coffee",
                category="Food",
                amount_rub=Decimal("100"),
            )
        )
    except ExpensesStorageError as error:
        assert "Invalid workbook headers" in str(error)
    else:
        raise AssertionError("Expected invalid headers to be rejected")

    assert path.read_bytes() == original


def test_repository_preserves_workbook_when_atomic_replace_fails(tmp_path, monkeypatch) -> None:
    path = tmp_path / "finances.xlsx"
    _create_workbook(path)
    original = path.read_bytes()

    def fail_replace(source, destination) -> None:
        raise PermissionError("read-only file")

    monkeypatch.setattr(expenses_repository.os, "replace", fail_replace)

    try:
        ExpensesRepository(path).add(
            ExpenseCreate(
                occurred_at=datetime.now(UTC),
                name="Coffee",
                category="Food",
                amount_rub=Decimal("100"),
            )
        )
    except ExpensesStorageError as error:
        assert str(error) == "Unable to save workbook: read-only file"
    else:
        raise AssertionError("Expected save failure")

    assert path.read_bytes() == original
    assert list(tmp_path.glob(".finances-*.xlsx")) == []


def test_expenses_api_paginates_and_creates_expense(monkeypatch, tmp_path) -> None:
    path = tmp_path / "finances.xlsx"
    _create_workbook(
        path,
        [
            _row("00000000-0000-0000-0000-000000000001", "2026-09-21T10:00:00+06:00"),
            _row("00000000-0000-0000-0000-000000000002", "2026-09-22T10:00:00+06:00"),
        ],
    )
    monkeypatch.setattr(service, "expenses", ExpensesRepository(path))
    client = TestClient(app)

    first_page = client.get("/expenses", params={"page_size": 1})
    assert first_page.status_code == 200
    assert first_page.json()["items"][0]["id"].endswith("2")
    cursor = first_page.json()["next_cursor"]
    second_page = client.get("/expenses", params={"page_size": 1, "cursor": cursor})
    assert second_page.status_code == 200
    assert second_page.json()["items"][0]["id"].endswith("1")
    assert second_page.json()["next_cursor"] is None

    created = client.post(
        "/expenses",
        json={
            "occurred_at": "2026-09-23T14:30:00+06:00",
            "name": "Gym",
            "category": "Sport",
            "amount_rub": "2500.00",
            "tags": ["health"],
        },
    )
    assert created.status_code == 201
    assert created.json()["name"] == "Gym"
    assert created.json()["amount_rub"] == "2500.00"


def test_expenses_api_returns_structured_errors(monkeypatch, tmp_path) -> None:
    missing_path = tmp_path / "missing.xlsx"
    monkeypatch.setattr(service, "expenses", ExpensesRepository(missing_path))
    client = TestClient(app)

    missing = client.get("/expenses")
    assert missing.status_code == 404
    assert missing.json()["detail"]["code"] == "finances_file_not_found"

    invalid_cursor = client.get("/expenses", params={"cursor": "***"})
    assert invalid_cursor.status_code == 422
    assert invalid_cursor.json()["detail"] == {
        "code": "invalid_cursor",
        "message": "Expense cursor is invalid",
    }

    invalid_page_size = client.get("/expenses", params={"page_size": 101})
    assert invalid_page_size.status_code == 422
    assert invalid_page_size.json()["detail"]["code"] == "invalid_expense_request"


def test_expense_request_validation_rejects_invalid_values(tmp_path, monkeypatch) -> None:
    path = tmp_path / "finances.xlsx"
    _create_workbook(path)
    monkeypatch.setattr(service, "expenses", ExpensesRepository(path))
    client = TestClient(app)
    valid = {
        "occurred_at": "2026-09-23T14:30:00+06:00",
        "name": "Coffee",
        "category": "Food",
        "amount_rub": "100.00",
    }

    invalid_payloads = [
        {**valid, "occurred_at": "2026-09-23T14:30:00"},
        {**valid, "amount_rub": "0"},
        {**valid, "amount_rub": "10.001"},
        {**valid, "name": "   "},
        {**valid, "tags": [""]},
        {**valid, "unexpected": True},
    ]

    for payload in invalid_payloads:
        response = client.post("/expenses", json=payload)
        assert response.status_code == 422
        assert response.json()["detail"]["code"] == "invalid_expense_request"


def test_expenses_api_returns_storage_error_for_malformed_row(monkeypatch, tmp_path) -> None:
    path = tmp_path / "finances.xlsx"
    _create_workbook(
        path,
        [
            _row(
                "00000000-0000-0000-0000-000000000001",
                "2026-09-22T10:00:00+06:00",
                tags="not-json",
            )
        ],
    )
    monkeypatch.setattr(service, "expenses", ExpensesRepository(path))

    response = TestClient(app).get("/expenses")

    assert response.status_code == 500
    assert response.json()["detail"] == {
        "code": "finances_storage_error",
        "message": "Invalid Tags JSON in row 2: Expecting value",
    }


def test_expenses_api_returns_storage_error_for_corrupt_workbook(monkeypatch, tmp_path) -> None:
    path = tmp_path / "finances.xlsx"
    path.write_text("not an xlsx file", encoding="utf-8")
    monkeypatch.setattr(service, "expenses", ExpensesRepository(path))

    response = TestClient(app).get("/expenses")

    assert response.status_code == 500
    assert response.json()["detail"]["code"] == "finances_storage_error"
    assert "Unable to open finances workbook" in response.json()["detail"]["message"]
